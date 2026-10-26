"""
Test Data Analyzer & Automatic Report Generator
-----------------------------------------------
Takes a raw measurement log (CSV) from an endurance / durability test and produces,
in one run:
  * an Excel workbook  - summary, anomaly list, highlighted data, native Excel charts
  * a PDF report       - KPIs, charts with limits and anomaly markers, findings table
  * PNG charts         - ready to paste into presentations

Detection:
  * Limit violations   - value outside the spec limits defined in SIGNALS
  * Statistical spikes - value far from the local (rolling) median, even if inside limits
  * Drift              - rolling average moving away from the start-of-test baseline
                         (early warning, before a limit is reached)
  * Data gaps          - missing cycle numbers (logger dropouts)

Usage:
    python generate_report.py sample_data/endurance_test_log.csv -o output
"""
from __future__ import annotations

import argparse
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

# --- Configuration: one line per measured signal --------------------------------
SIGNALS = {
    # column              label                  unit   low    high   drift % vs baseline
    "peak_pressure_bar": ("Peak pressure",       "bar", 3.70,  4.30,  3.0),
    "cycle_time_s":      ("Cycle time",          "s",   11.0,  14.0,  5.0),
    "current_A":         ("Current consumption", "A",   1.90,  2.40,  5.0),
    "temperature_C":     ("Temperature",         "°C",  15.0,  35.0,  10.0),
}
BASELINE_CYCLES = 100   # first N cycles define "normal"
ROLL = 25               # rolling window (cycles)
SPIKE_SIGMA = 4         # spike = this many baseline std-devs away from local median


def analyse(df: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    findings = []
    drift_start = {}
    for col, (label, unit, lo, hi, drift_pct) in SIGNALS.items():
        s = df[col]
        # 1) limit violations
        for idx in df.index[(s < lo) | (s > hi)]:
            findings.append((df.at[idx, "cycle"], label, "Limit violation",
                             f"{s[idx]:.3f} {unit} (limits {lo}–{hi})", "HIGH"))
        # 2) statistical spikes: far from the local median, measured in baseline std-devs
        med = s.rolling(ROLL, center=True, min_periods=5).median()
        b = s.iloc[:BASELINE_CYCLES]
        sigma = (b - b.median()).abs().median() * 1.4826   # robust std-dev (ignores early outliers)
        z = (s - med) / sigma
        for idx in df.index[(z.abs() > SPIKE_SIGMA) & s.between(lo, hi)]:
            findings.append((df.at[idx, "cycle"], label, "Spike",
                             f"{s[idx]:.3f} {unit} (local median {med[idx]:.3f})", "MEDIUM"))
        # 3) drift vs baseline
        base = s.iloc[:BASELINE_CYCLES].mean()
        roll = s.rolling(50, min_periods=50).mean()
        dev = (roll - base).abs() / base * 100
        drifting = df.index[dev > drift_pct]
        if len(drifting):
            first = drifting[0]
            drift_start[col] = df.at[first, "cycle"]
            findings.append((df.at[first, "cycle"], label, "Drift",
                             f"avg {roll[first]:.3f} {unit} vs baseline {base:.3f} (>{drift_pct}%)", "MEDIUM"))
    # 4) gaps in cycle numbering
    expected = set(range(int(df["cycle"].min()), int(df["cycle"].max()) + 1))
    missing = sorted(expected - set(df["cycle"]))
    if missing:
        findings.append((missing[0], "Data logger", "Data gap",
                         f"{len(missing)} cycle(s) missing: {missing[0]}–{missing[-1]}", "LOW"))

    f = pd.DataFrame(findings, columns=["Cycle", "Signal", "Type", "Detail", "Severity"])
    f["sev_order"] = f["Severity"].map({"HIGH": 0, "MEDIUM": 1, "LOW": 2})
    f = f.sort_values(["sev_order", "Cycle"]).drop(columns="sev_order").reset_index(drop=True)

    kpis = {
        "Cycles recorded": len(df),
        "Test duration": "{:.0f} h {:02.0f} min".format(*divmod((df["timestamp"].iloc[-1] - df["timestamp"].iloc[0]).total_seconds() / 60, 60)),
        "Limit violations": int((f["Type"] == "Limit violation").sum()),
        "Spikes": int((f["Type"] == "Spike").sum()),
        "Drifting signals": len(drift_start),
        "Missing cycles": len(missing),
        "Verdict": "FAIL" if (f["Severity"] == "HIGH").any() else ("WARNING" if len(f) else "PASS"),
    }
    return f, {"kpis": kpis, "drift_start": drift_start}


def make_charts(df, findings, meta, out: Path) -> list[Path]:
    paths = []
    for col, (label, unit, lo, hi, _) in SIGNALS.items():
        fig, ax = plt.subplots(figsize=(10, 3.4), dpi=130)
        ax.plot(df["cycle"], df[col], lw=0.9, color="#2F5D8A", label=label)
        ax.plot(df["cycle"], df[col].rolling(50, min_periods=50).mean(), lw=2, color="#E39B2D", label="50-cycle average")
        ax.axhline(lo, color="#C0392B", ls="--", lw=1, label="Spec limits")
        ax.axhline(hi, color="#C0392B", ls="--", lw=1)
        hits = findings[(findings["Signal"] == label) & findings["Type"].isin(["Limit violation", "Spike"])]
        pts = df[df["cycle"].isin(hits["Cycle"])]
        ax.scatter(pts["cycle"], pts[col], s=45, color="#C0392B", zorder=5, label="Anomaly")
        if col in meta["drift_start"]:
            ax.axvline(meta["drift_start"][col], color="#E39B2D", ls=":", lw=1.5, label="Drift detected")
        ax.set_title(f"{label} [{unit}]", loc="left", fontsize=11, fontweight="bold")
        ax.set_xlabel("Cycle")
        ax.grid(alpha=0.25)
        ax.spines[["top", "right"]].set_visible(False)
        ax.legend(fontsize=7.5, ncol=5, loc="upper center", bbox_to_anchor=(0.5, -0.22), frameon=False)
        fig.tight_layout()
        p = out / f"chart_{col}.png"
        fig.savefig(p)
        plt.close(fig)
        paths.append(p)
    return paths


def write_excel(df, findings, meta, path: Path):
    head_fill, head_font = PatternFill("solid", fgColor="1F3A5F"), Font(bold=True, color="FFFFFF")
    red = PatternFill("solid", fgColor="F8D7DA")
    kpi = pd.DataFrame(list(meta["kpis"].items()), columns=["Metric", "Value"])
    data = df.copy()
    data["timestamp"] = data["timestamp"].dt.strftime("%Y-%m-%d %H:%M:%S")
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        kpi.to_excel(w, sheet_name="Summary", index=False)
        findings.to_excel(w, sheet_name="Anomalies", index=False)
        data.to_excel(w, sheet_name="Data", index=False)
        for ws in w.sheets.values():
            for c in ws[1]:
                c.fill, c.font = head_fill, head_font
            for i, col in enumerate(ws.iter_cols(), 1):
                ws.column_dimensions[get_column_letter(i)].width = min(
                    max(len(str(c.value)) if c.value is not None else 0 for c in col) + 3, 60)
            ws.freeze_panes = "A2"
        # verdict colour
        v = w.sheets["Summary"]["B8"]
        v.font = Font(bold=True, color={"PASS": "1F7A4D", "WARNING": "B7791F", "FAIL": "B42318"}[v.value])
        # highlight out-of-limit cells in the data sheet
        ws = w.sheets["Data"]
        cols = list(data.columns)
        for col, (_, _, lo, hi, _) in SIGNALS.items():
            ci = cols.index(col) + 1
            for r in range(2, ws.max_row + 1):
                val = ws.cell(r, ci).value
                if val is not None and not (lo <= val <= hi):
                    ws.cell(r, ci).fill = red
        # native Excel line charts (stay editable for the client)
        cs = w.book.create_sheet("Charts")
        for n, col in enumerate(SIGNALS):
            ci = cols.index(col) + 1
            ch = LineChart()
            ch.title = f"{SIGNALS[col][0]} [{SIGNALS[col][1]}]"
            ch.height, ch.width = 7, 24
            ch.legend = None
            ch.add_data(Reference(ws, min_col=ci, min_row=1, max_row=ws.max_row), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
            ch.series[0].graphicalProperties.line.width = 12000
            cs.add_chart(ch, f"A{1 + n * 15}")


def write_pdf(findings, meta, charts, path: Path, source: str):
    st = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm,
                            topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    k = meta["kpis"]
    verdict_color = {"PASS": "#1F7A4D", "WARNING": "#B7791F", "FAIL": "#B42318"}[k["Verdict"]]
    story = [Paragraph("Endurance Test Report", st["Title"]),
             Paragraph(f"Source file: {source} · generated automatically", st["Normal"]), Spacer(1, 12),
             Paragraph(f'<font size=16 color="{verdict_color}"><b>Verdict: {k["Verdict"]}</b></font>', st["Normal"]),
             Spacer(1, 12)]
    kt = Table([[a, str(b)] for a, b in k.items() if a != "Verdict"], colWidths=[7 * cm, 6 * cm])
    kt.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF2F7")),
                            ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold")]))
    story += [kt, Spacer(1, 16), Paragraph("Findings", st["Heading2"])]
    rows = [["Cycle", "Signal", "Type", "Detail", "Severity"]] + findings.astype(str).values.tolist()
    ft = Table(rows, colWidths=[1.4 * cm, 3.6 * cm, 2.8 * cm, 7 * cm, 2 * cm], repeatRows=1)
    style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
             ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTSIZE", (0, 0), (-1, -1), 8),
             ("GRID", (0, 0), (-1, -1), 0.3, colors.grey)]
    for i, sev in enumerate(findings["Severity"], start=1):
        c = {"HIGH": "#F8D7DA", "MEDIUM": "#FFF1D6", "LOW": "#EEF2F7"}[sev]
        style.append(("BACKGROUND", (4, i), (4, i), colors.HexColor(c)))
    ft.setStyle(TableStyle(style))
    story += [ft, PageBreak(), Paragraph("Signal trends", st["Heading2"])]
    for p in charts:
        story += [Image(str(p), width=17.4 * cm, height=17.4 * cm * 3.4 / 10), Spacer(1, 6)]
    doc.build(story)


def main():
    ap = argparse.ArgumentParser(description="Analyse a test log and generate Excel + PDF reports.")
    ap.add_argument("csv", type=Path)
    ap.add_argument("-o", "--out", type=Path, default=Path("output"))
    args = ap.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)

    df = pd.read_csv(args.csv, parse_dates=["timestamp"])
    findings, meta = analyse(df)
    charts = make_charts(df, findings, meta, args.out)
    write_excel(df, findings, meta, args.out / "test_report.xlsx")
    write_pdf(findings, meta, charts, args.out / "test_report.pdf", args.csv.name)

    for kk, vv in meta["kpis"].items():
        print(f"  {kk:<18} {vv}")
    print(f"\n{len(findings)} findings. Reports saved in {args.out}/ (xlsx, pdf, png)")


if __name__ == "__main__":
    main()
